"""Build full, timestamped transcripts without summarizing or silent omissions."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from transcribe_local import digest, save, signature


def clock(seconds):
    s = int(seconds)
    return f'{s//3600:02d}:{s//60%60:02d}:{s%60:02d}'


def validate_transcript(t):
    if not t.get('complete') or not t.get('segments') or not t.get('blocks'):
        raise ValueError('incomplete_transcript')
    end = 0
    for b in t['blocks']:
        if b['status'] != 'complete' or abs(b['start']-end) > .001 or b['end'] <= end:
            raise ValueError('audio_block_gap_or_repeat')
        end = b['end']
    if abs(end-t['duration']) > .001:
        raise ValueError('audio_coverage_incomplete')
    previous_start = 0
    for s in t['segments']:
        if not 0 <= s['start'] <= s['end'] <= t['duration'] or s['start'] < previous_start:
            raise ValueError('invalid_segment_order')
        if not s['text'].strip():
            raise ValueError('empty_segment')
        previous_start = s['start']


def groups(segments):
    current = []
    for s in segments:
        if current and s['start']-current[0]['start'] >= 60:
            yield current
            current = []
        current.append(s)
    if current:
        yield current


def build(manifest_path, out):
    manifest = json.loads(manifest_path.read_text())
    entries = manifest['replays']
    if not entries:
        raise ValueError('empty_manifest')
    records = []
    ids = set()
    for e in entries:
        ident = (e['account_id'], e['replay_id'])
        if ident in ids:
            raise ValueError('duplicate_replay')
        ids.add(ident)
        p = Path(e['transcript'])
        t = json.loads(p.read_text())
        validate_transcript(t)
        if t['identity'] != dict(account=e['account_id'], replay=e['replay_id']):
            raise ValueError('transcript_identity_mismatch')
        records.append((e, t, digest(p)))
    sig = signature({'manifest': manifest, 'transcripts': [r[2] for r in records], 'builder': digest(__file__)})
    reportpath = out.with_suffix('.verification.json')
    if out.exists():
        report = json.loads(reportpath.read_text()) if reportpath.exists() else {}
        if report.get('signature') == sig and report.get('docx_sha256') == digest(out):
            return report
        raise ValueError('existing_word_preserved')
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name, size in [('Normal', 11), ('Title', 22), ('Heading 1', 16), ('Heading 2', 12)]:
        style = d.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = 'PingFang SC', Pt(size), RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'PingFang SC')
        style.paragraph_format.space_after = Pt(7)
    d.styles['Normal'].paragraph_format.line_spacing = 1.25
    body_style = d.styles.add_style('TranscriptBody', 1)
    body_style.base_style = d.styles['Normal']
    # Strip template theme fonts and inherited title rules before rendering.
    for element in list(d.styles.element.iter()):
        if element.tag == qn('w:pBdr'):
            element.getparent().remove(element)
        elif element.tag == qn('w:rFonts'):
            for attr in list(element.attrib):
                if 'theme' in attr.rsplit('}', 1)[-1].lower():
                    del element.attrib[attr]
            for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                element.set(qn('w:'+slot), 'PingFang SC')
    kind = '课程' if manifest.get('kind') == 'course' else '直播回放'
    d.add_paragraph(manifest['account_name']+' '+kind+'文字稿', 'Title')
    d.add_paragraph(f'本文件收录 {len(entries)} 节{kind}的完整机器转写，按清单与原始音频顺序编排。时间戳为各节内播放位置，正文未作摘要。')
    if not manifest.get('catalog_complete'):
        d.add_paragraph('本文件仅对应以下场次，不代表该账号全部历史回放已收齐。')
    d.add_paragraph('文字由本地语音识别生成，未逐字人工校听。人名、数字及专有名词可能有识别误差；低置信度片段在对应场次列明。')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    expected = []
    for index, (e, t, _) in enumerate(records):
        if index:
            d.add_page_break()
        title = re.sub(r'[^\w\s\u3400-\u9fff]', '', e['title'])
        d.add_paragraph(title, 'Heading 1')
        d.add_paragraph('原题：'+e['title'])
        d.add_paragraph('来源：'+e['source_url'])
        d.add_paragraph(f"日期：{e.get('date') or '页面未核实'}　时长：{clock(t['duration'])}")
        issues = [x for x in t.get('quality_issues', []) if x['type']=='low_confidence']
        if issues:
            d.add_paragraph('听辨待核位置：'+'、'.join(clock(t['segments'][x['segment']]['start']) for x in issues))
        for chunk in groups(t['segments']):
            stamp = f"[{clock(chunk[0]['start'])}–{clock(chunk[-1]['end'])}] "
            text = ''.join(s['text'] for s in chunk)
            p = d.add_paragraph(style='TranscriptBody')
            p.add_run(stamp).bold = True
            p.add_run(text)
            expected.append(stamp+text)
        if e.get('source_text'):
            d.add_paragraph('用户提供的课程讲义原文', 'Heading 2')
            for text in e['source_text']:
                d.add_paragraph(text, 'TranscriptBody')
                expected.append(text)
    for paragraph in d.paragraphs:
        for run in paragraph.runs:
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn('w:'+slot), 'PingFang SC')
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix('.partial.docx')
    if tmp.exists():
        raise ValueError('partial_word_preserved')
    d.save(tmp)
    actual = [p.text for p in Document(tmp).paragraphs if p.style.name=='TranscriptBody']
    if actual != expected:
        raise ValueError('word_body_coverage_mismatch')
    tmp.replace(out)
    report = dict(signature=sig, docx_sha256=digest(out), replay_count=len(records),
                  paragraph_count=len(expected), body_exact_match=True,
                  catalog_complete=bool(manifest.get('catalog_complete')), render_review_complete=False)
    save(reportpath, report)
    return report


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('manifest', type=Path); p.add_argument('output', type=Path)
    a = p.parse_args()
    print(json.dumps(build(a.manifest, a.output), ensure_ascii=False))
